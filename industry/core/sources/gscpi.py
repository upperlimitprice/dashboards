"""뉴욕연은 글로벌 공급망 압력지수(GSCPI) 엑셀."""
import io
from datetime import datetime

from ..http import get

URL = "https://www.newyorkfed.org/medialibrary/research/interactives/gscpi/downloads/gscpi_data.xlsx"


def parse(xlsx_bytes):
    if xlsx_bytes[:4] == b"\xd0\xcf\x11\xe0":  # 확장자는 .xlsx지만 실제로는 구형 .xls(OLE)로 배포됨
        return _parse_xls(xlsx_bytes)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            if len(r) >= 2 and isinstance(r[0], datetime) and isinstance(r[1], (int, float)):
                out[r[0].strftime("%Y-%m")] = float(r[1])
        if out:
            break
    wb.close()
    return out


def _parse_xls(b):
    import xlrd
    wb = xlrd.open_workbook(file_contents=b)
    out = {}
    for sh in wb.sheets():
        for i in range(sh.nrows):
            r = sh.row_values(i)
            if len(r) >= 2 and isinstance(r[0], float) and r[0] > 20000 and isinstance(r[1], float):
                d = xlrd.xldate_as_datetime(r[0], wb.datemode)
                out[d.strftime("%Y-%m")] = r[1]
            elif len(r) >= 2 and isinstance(r[0], str) and isinstance(r[1], float):
                try:
                    out[datetime.strptime(r[0].strip(), "%d-%b-%Y").strftime("%Y-%m")] = r[1]  # '31-Jan-1998' 문자열 날짜
                except ValueError:
                    pass
        if out:
            break
    return out


def series():
    return parse(get(URL, timeout=90, cache_hours=12))
