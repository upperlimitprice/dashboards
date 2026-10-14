"""EIA-860M 월간 발전기 인벤토리 엑셀 (키 불필요).

https://www.eia.gov/electricity/data/eia860m/ — 최신 파일은 xls/{month}_generator{year}.xlsx.
시트: Operating / Planned / Retired. 헤더 행은 'Nameplate Capacity (MW)' 셀로 탐지.
"""
import calendar
import io
from datetime import date

from ..http import NotFound, get

URL = "https://www.eia.gov/electricity/data/eia860m/xls/{m}_generator{y}.xlsx"
SHEETS = ("Operating", "Planned", "Retired")

TECH_GROUPS = [   # (부분문자열, 그룹) — 위에서부터 첫 매치
    ("Solar", "태양광"), ("Batter", "ESS"), ("Wind", "풍력"), ("Natural Gas", "가스"),
    ("Coal", "석탄"), ("Nuclear", "원전"), ("Hydro", "수력"), ("Petroleum", "석유"),
]
GROUP_ORDER = ["태양광", "ESS", "풍력", "가스", "원전", "석탄", "수력", "석유", "기타"]

# Planned 시트 Status 코드: (V)>50% 공사, (U)≤50% 공사, (T) 인허가 완료, (L) 인허가 중, (P) 계획, (OT) 기타
UNDER_CONSTRUCTION = ("V", "U")


def tech_group(tech):
    tech = tech or ""
    for needle, g in TECH_GROUPS:
        if needle.lower() in tech.lower():
            return g
    return "기타"


def status_code(status):
    s = (status or "").strip()
    return s[1:s.index(")")] if s.startswith("(") and ")" in s else s


def latest(today=None, lookback=8):
    """가장 최근 공개 파일 (bytes, 'YYYY-MM'). 보통 2개월 전 데이터가 최신."""
    today = today or date.today()
    y, m = today.year, today.month
    last_err = None
    for _ in range(lookback):
        url = URL.format(m=calendar.month_name[m].lower(), y=y)
        try:
            return get(url, timeout=180, cache_hours=24), f"{y:04d}-{m:02d}"
        except NotFound as e:
            last_err = e
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    raise RuntimeError(f"EIA-860M 파일을 찾지 못함: {last_err}")


def parse(xlsx_bytes):
    """{시트명: [ {헤더: 값}, ... ]}"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    out = {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            continue
        rows = wb[name].iter_rows(values_only=True)
        header = None
        for _ in range(10):
            r = next(rows, None)
            if r is None:
                break
            cells = [str(c).strip() if c is not None else "" for c in r]
            if "Nameplate Capacity (MW)" in cells:
                header = cells
                break
        if not header:
            continue
        recs = []
        for r in rows:
            if not r or all(c is None for c in r):
                continue
            rec = dict(zip(header, r))
            if not isinstance(rec.get("Nameplate Capacity (MW)"), (int, float)):
                continue      # 각주 행
            recs.append(rec)
        out[name] = recs
    wb.close()
    return out


def summarize(sheets):
    """대시보드용 요약 — 원본 2만여 행 대신 집계만 저장."""
    op, pl, rt = (sheets.get(s, []) for s in SHEETS)
    mw = "Nameplate Capacity (MW)"

    def bucket(recs, key_fn):
        acc = {}
        for r in recs:
            k = key_fn(r)
            if k is None:
                continue
            g = tech_group(r.get("Technology"))
            acc.setdefault(str(k), {}).setdefault(g, 0.0)
            acc[str(k)][g] += float(r[mw])
        return acc

    def year(field):
        def f(r):
            v = r.get(field)
            try:
                return int(v)
            except (TypeError, ValueError):
                return None
        return f

    planned_uc = [r for r in pl if status_code(r.get("Status")) in UNDER_CONSTRUCTION]
    state_acc = {}
    for r in pl:
        st = r.get("Plant State") or "?"
        d = state_acc.setdefault(st, {"all": 0.0, "uc": 0.0, "solar_ess": 0.0})
        v = float(r[mw])
        d["all"] += v
        if status_code(r.get("Status")) in UNDER_CONSTRUCTION:
            d["uc"] += v
        if tech_group(r.get("Technology")) in ("태양광", "ESS"):
            d["solar_ess"] += v
    return {
        "operating_by_tech": bucket(op, lambda r: "total").get("total", {}),
        "additions_by_year": bucket(op, year("Operating Year")),
        "planned_by_year": bucket(pl, year("Planned Operation Year")),
        "planned_uc_by_year": bucket(planned_uc, year("Planned Operation Year")),
        "planned_monthly": bucket(pl, lambda r: (
            f"{int(r['Planned Operation Year'])}-{int(r['Planned Operation Month']):02d}"
            if r.get("Planned Operation Year") and r.get("Planned Operation Month") else None)),
        "retired_by_year": bucket(rt, year("Retirement Year")),
        "planned_retire_by_year": bucket(op, year("Planned Retirement Year")),
        "planned_by_state": state_acc,
    }
