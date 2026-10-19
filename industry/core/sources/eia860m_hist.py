"""EIA-860M 월별 스냅샷 이력 — 발전기 단위(Plant ID + Generator ID) 추적용 최소 필드 추출.

최신 파일: https://www.eia.gov/electricity/data/eia860m/xls/{month}_generator{year}.xlsx
과거 파일: https://www.eia.gov/electricity/data/eia860m/archive/xls/{month}_generator{year}.xlsx
미공개 월은 404 대신 200 HTML 을 돌려주므로 b[:2] == b"PK" 로 xlsx 여부를 확인한다.

스냅샷 토큰(발전기·월별 상태 1개 문자열):
  "2608V"  Planned 시트, 예정 가동 2026-08, 상태 코드 V
  "O2607"  Operating 시트, 가동 개시 2026-07
  "C"      Canceled or Postponed 시트
  "R"      Retired 시트
  ""       어느 시트에도 없음
"""
import calendar
import io

from ..http import NotFound, get

CUR = "https://www.eia.gov/electricity/data/eia860m/xls/{m}_generator{y}.xlsx"
ARCH = "https://www.eia.gov/electricity/data/eia860m/archive/xls/{m}_generator{y}.xlsx"

# 추적 기술: 860M 'Technology' 값 → 코드
TECH = {"Solar Photovoltaic": "S", "Batteries": "B", "Onshore Wind Turbine": "W",
        "Offshore Wind Turbine": "W", "Nuclear": "N"}

SHEET_TAG = {"Operating": "O", "Planned": "P", "Canceled or Postponed": "C", "Retired": "R"}


def _urls(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    name = calendar.month_name[m].lower()
    return [CUR.format(m=name, y=y), ARCH.format(m=name, y=y)]


def download(ym, cache_hours=24 * 7):
    """해당 월 파일 bytes. 최신 월은 xls/, 이전 월은 archive/xls/ 에 있다. 없으면 None.
    미공개 월 탐색(소프트 404 HTML 도 캐시됨)에는 짧은 cache_hours 를 쓸 것."""
    for url in _urls(ym):
        try:
            b = get(url, timeout=240, cache_hours=cache_hours)
        except NotFound:
            continue
        if b[:2] == b"PK":
            return b
    return None


def _int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _code(status):
    s = (status or "").strip()
    return s[1:s.index(")")] if s.startswith("(") and ")" in s else (s[:2] or "?")


PRIORITY = {"P": 0, "O": 1, "C": 2, "R": 3}     # 같은 키가 여러 시트에 있으면 Planned > Operating > 취소 > 폐지


def parse(xlsx_bytes):
    """→ (gens, op_mw)
    gens  = {key: (tech, state, mw, token)}  추적 기술만. key = 'PlantID-GenID'
    op_mw = {tech: {"US": MW, "TX": MW, "CA": MW, ...}}  Operating 시트 기술·주별 가동 용량
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    gens, op_mw = {}, {}
    for sheet, tag in SHEET_TAG.items():
        if sheet not in wb.sheetnames:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        idx = None
        for _ in range(10):
            r = next(rows, None)
            if r is None:
                break
            cells = [str(c).strip() if c is not None else "" for c in r]
            if "Nameplate Capacity (MW)" in cells:
                idx = {c: i for i, c in enumerate(cells)}
                break
        if not idx:
            continue

        def g(r, c):
            i = idx.get(c)
            return r[i] if i is not None and i < len(r) else None
        for r in rows:
            if not r:
                continue
            t = TECH.get(str(g(r, "Technology") or "").strip())
            mw = g(r, "Nameplate Capacity (MW)")
            pid, gid = g(r, "Plant ID"), g(r, "Generator ID")
            if not t or not isinstance(mw, (int, float)) or pid in (None, "") or gid in (None, ""):
                continue
            st = str(g(r, "Plant State") or "?").strip()
            if tag == "P":
                y, m = _int(g(r, "Planned Operation Year")), _int(g(r, "Planned Operation Month"))
                if not y or not m:
                    continue
                tok = f"{y % 100:02d}{m:02d}{_code(g(r, 'Status'))}"
            elif tag == "O":
                d = op_mw.setdefault(t, {})
                d["US"] = d.get("US", 0) + float(mw)
                d[st] = d.get(st, 0) + float(mw)
                y, m = _int(g(r, "Operating Year")), _int(g(r, "Operating Month"))
                tok = f"O{y % 100:02d}{m:02d}" if y and m else "O"
            else:
                tok = tag
            key = f"{_int(pid) or pid}-{str(gid).strip()}"
            prev = gens.get(key)
            if prev and PRIORITY[_tag(prev[3])] <= PRIORITY[tag]:
                continue
            gens[key] = (t, st, round(float(mw), 1), tok)
    wb.close()
    return gens, {t: {k: round(v, 1) for k, v in d.items()} for t, d in op_mw.items()}


def _tag(tok):
    return tok[0] if tok and tok[0] in "OCR" else "P"
